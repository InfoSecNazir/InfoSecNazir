import json
from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q, Avg
from django.utils import timezone
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO
from .models import Student, Session, AttendanceLog, AuditLog, Badge, StudentBadge


def home(request):
    """Home page with system overview"""
    context = {
        'total_students': Student.objects.count(),
        'total_sessions': Session.objects.count(),
        'total_attendance': AttendanceLog.objects.count(),
        'recent_sessions': Session.objects.all()[:5],
    }
    return render(request, 'students/home.html', context)


def student_list(request):
    """List all students with their statistics"""
    students = Student.objects.all()
    
    # Apply filters
    search_query = request.GET.get('search', '')
    if search_query:
        students = students.filter(
            Q(name__icontains=search_query) |
            Q(university_id__icontains=search_query) |
            Q(email__icontains=search_query)
        )
    
    # Filter by attendance percentage
    attendance_filter = request.GET.get('attendance', '')
    if attendance_filter == 'low':
        students = students.filter(attendance_percentage__lt=50)
    elif attendance_filter == 'medium':
        students = students.filter(attendance_percentage__gte=50, attendance_percentage__lt=75)
    elif attendance_filter == 'high':
        students = students.filter(attendance_percentage__gte=75)
    
    context = {
        'students': students,
        'search_query': search_query,
        'attendance_filter': attendance_filter,
    }
    return render(request, 'students/student_list.html', context)


def student_detail(request, student_id):
    """Student detail page with attendance history"""
    student = get_object_or_404(Student, id=student_id)
    attendance_logs = AttendanceLog.objects.filter(student=student)
    badges = StudentBadge.objects.filter(student=student)
    
    context = {
        'student': student,
        'attendance_logs': attendance_logs,
        'badges': badges,
    }
    return render(request, 'students/student_detail.html', context)


def scanner_interface(request):
    """QR Code scanner interface"""
    sessions = Session.objects.filter(date=timezone.now().date())
    context = {
        'sessions': sessions,
        'today': timezone.now().date(),
    }
    return render(request, 'students/scanner.html', context)


@csrf_exempt
def process_qr_scan(request):
    """Process scanned QR code and record attendance"""
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            qr_data = data.get('qr_data', '')
            session_id = data.get('session_id')
            status = data.get('status', 'present')
            latitude = data.get('latitude')
            longitude = data.get('longitude')
            
            # Extract UUID from QR data
            if not qr_data.startswith('STUDENT:'):
                return JsonResponse({'success': False, 'error': 'Invalid QR code format'})
            
            uuid_str = qr_data.replace('STUDENT:', '').strip()
            
            # Find student
            try:
                student = Student.objects.get(uuid=uuid_str)
            except Student.DoesNotExist:
                return JsonResponse({'success': False, 'error': 'Student not found'})
            
            # Find session
            session = get_object_or_404(Session, id=session_id)
            
            # Check if attendance already recorded
            attendance, created = AttendanceLog.objects.get_or_create(
                student=student,
                session=session,
                defaults={
                    'status': status,
                    'scan_latitude': latitude,
                    'scan_longitude': longitude,
                    'recorded_by': request.user.username if request.user.is_authenticated else 'Anonymous',
                }
            )
            
            if not created:
                # Update existing record
                attendance.status = status
                attendance.scan_latitude = latitude
                attendance.scan_longitude = longitude
                attendance.save()
            
            # Update student attendance percentage
            total_sessions = Session.objects.filter(date__lte=timezone.now().date()).count()
            attended_sessions = AttendanceLog.objects.filter(
                student=student,
                status__in=['present', 'late']
            ).count()
            
            if total_sessions > 0:
                student.attendance_percentage = (attended_sessions / total_sessions) * 100
                student.save()
            
            # Create audit log
            AuditLog.objects.create(
                action_type='scan',
                user=request.user.username if request.user.is_authenticated else 'Anonymous',
                description=f"QR scanned for {student.name} - {session.title}",
                ip_address=request.META.get('REMOTE_ADDR')
            )
            
            return JsonResponse({
                'success': True,
                'message': f"Attendance recorded for {student.name}",
                'student': {
                    'name': student.name,
                    'university_id': student.university_id,
                    'status': status,
                }
            })
            
        except Exception as e:
            return JsonResponse({'success': False, 'error': str(e)})
    
    return JsonResponse({'success': False, 'error': 'Invalid request method'})


def dashboard(request):
    """Admin dashboard with statistics and analytics"""
    # Calculate statistics
    total_students = Student.objects.count()
    total_sessions = Session.objects.count()
    total_attendance = AttendanceLog.objects.count()
    
    # Students at risk (attendance < 50%)
    at_risk_students = Student.objects.filter(attendance_percentage__lt=50).count()
    
    # Recent attendance logs
    recent_logs = AttendanceLog.objects.select_related('student', 'session')[:10]
    
    # Attendance by session type
    session_stats = Session.objects.values('session_type').annotate(
        count=Count('id')
    )
    
    # Top performers
    top_students = Student.objects.order_by('-attendance_percentage')[:10]
    
    context = {
        'total_students': total_students,
        'total_sessions': total_sessions,
        'total_attendance': total_attendance,
        'at_risk_students': at_risk_students,
        'recent_logs': recent_logs,
        'session_stats': session_stats,
        'top_students': top_students,
    }
    return render(request, 'students/dashboard.html', context)


def export_attendance_excel(request):
    """Export attendance data to Excel"""
    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Report"
    
    # Style definitions
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    # Headers
    headers = ['University ID', 'Student Name', 'Email', 'Attendance %', 'Total Sessions', 'Present', 'Absent', 'Late']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # Data
    students = Student.objects.all()
    for row, student in enumerate(students, 2):
        present_count = AttendanceLog.objects.filter(student=student, status='present').count()
        absent_count = AttendanceLog.objects.filter(student=student, status='absent').count()
        late_count = AttendanceLog.objects.filter(student=student, status='late').count()
        total = present_count + absent_count + late_count
        
        ws.cell(row=row, column=1, value=student.university_id)
        ws.cell(row=row, column=2, value=student.name)
        ws.cell(row=row, column=3, value=student.email)
        ws.cell(row=row, column=4, value=float(student.attendance_percentage))
        ws.cell(row=row, column=5, value=total)
        ws.cell(row=row, column=6, value=present_count)
        ws.cell(row=row, column=7, value=absent_count)
        ws.cell(row=row, column=8, value=late_count)
    
    # Adjust column widths
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
    
    # Save to BytesIO
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    # Create response
    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=attendance_report_{timezone.now().strftime("%Y%m%d")}.xlsx'
    
    # Create audit log
    AuditLog.objects.create(
        action_type='export',
        user=request.user.username if request.user.is_authenticated else 'Anonymous',
        description=f"Exported attendance report to Excel",
        ip_address=request.META.get('REMOTE_ADDR')
    )
    
    return response


def session_list(request):
    """List all sessions"""
    sessions = Session.objects.all()
    
    # Filter by session type
    session_type = request.GET.get('type', '')
    if session_type:
        sessions = sessions.filter(session_type=session_type)
    
    # Filter by date range
    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')
    if date_from:
        sessions = sessions.filter(date__gte=date_from)
    if date_to:
        sessions = sessions.filter(date__lte=date_to)
    
    context = {
        'sessions': sessions,
        'session_type': session_type,
        'date_from': date_from,
        'date_to': date_to,
    }
    return render(request, 'students/session_list.html', context)


def session_detail(request, session_id):
    """Session detail with attendance records"""
    session = get_object_or_404(Session, id=session_id)
    attendance_logs = AttendanceLog.objects.filter(session=session).select_related('student')
    
    # Statistics
    total_students = Student.objects.count()
    present_count = attendance_logs.filter(status='present').count()
    absent_count = attendance_logs.filter(status='absent').count()
    late_count = attendance_logs.filter(status='late').count()
    
    context = {
        'session': session,
        'attendance_logs': attendance_logs,
        'total_students': total_students,
        'present_count': present_count,
        'absent_count': absent_count,
        'late_count': late_count,
    }
    return render(request, 'students/session_detail.html', context)
